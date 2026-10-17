"""End-to-end tests for the Presentation Studio workflow.

These walk the *complete* scenario a user performs -- create a project, discuss a
technical subject, draft and approve an outline, generate the slides in batches,
edit and revise them, approve the batch, reorder, close and reopen, approve the
whole deck, and export -- and then reopen the written ``.pptx`` with python-pptx
and assert its real structure (slide count, order, titles, content, formatting,
and notes).

Everything runs through the real :class:`gui.presentation.PresentationService`
and the real ``cli.presentation`` entry point, so the tests exercise the code the
window and the CLI actually call. No model and no network are required: the
"assistant" is supplied either as a deterministic callable or as text, so the
tests are fast and reproducible.

The workflow tests run on a base install (the stdlib writer). The tests that
reopen a deck with python-pptx skip when the optional ``[pptx]`` extra is absent.
"""

from __future__ import annotations

import json
import zipfile
from collections.abc import Callable
from pathlib import Path

import pytest

from core.presentation import (
    CONTENT_EQUATION,
    CONTENT_IMAGE,
    CONTENT_TABLE,
    OutlineSection,
    STATE_CONTENT_APPROVED,
    STATE_EXPORTED,
    STATE_SLIDE_REVIEW,
)
from gui.presentation import PresentationService, SlideSpec

try:  # the optional extra; every structural assertion below needs it
    import pptx as _pptx  # noqa: F401

    PPTX_AVAILABLE = True
except Exception:  # noqa: BLE001 - absence is a legitimate, reported state
    PPTX_AVAILABLE = False

requires_pptx = pytest.mark.skipif(
    not PPTX_AVAILABLE, reason="the optional python-pptx extra is not installed"
)


# --------------------------------------------------------------------------- #
# A deterministic assistant (no model, no network)
# --------------------------------------------------------------------------- #

#: The outline the scripted assistant proposes: two sections, with subsections.
#: A subsection is an indented (or sub-numbered) line under its section.
OUTLINE = (
    "Reservoir Characterisation\n"
    "    - Porosity and storage\n"
    "    - Permeability and flow\n"
    "Well Logs\n"
    "    - Gamma ray\n"
    "    - Resistivity"
)

#: A prose answer, so the drafting path is exercised with real sentences.
PROSE = (
    "Porosity is the fraction of the bulk volume that is void space. "
    "Effective porosity counts only the pores that are connected to each other. "
    "Permeability is the ability of the rock to transmit a fluid. "
    "Darcy's law relates flow rate to permeability and pressure gradient."
)


class ScriptedAssistant:
    """A deterministic stand-in for the project's assistant."""

    def __init__(self, answer: str = PROSE) -> None:
        self.answer = answer
        self.questions: list[str] = []

    def __call__(self, question: str) -> str:
        self.questions.append(question)
        return self.answer


def _service(
    tmp_path: Path, *, assistant: Callable[[str], str] | None = None
) -> PresentationService:
    return PresentationService(
        root=tmp_path / "store",
        output_dir=tmp_path / "out",
        ask=assistant if assistant is not None else ScriptedAssistant(),
    )


def _slide_text(slide) -> str:
    """Every run of text on a slide, including table cells, in shape order.

    The engine draws a deck as real text boxes (and real tables) on a blank
    layout rather than filling placeholders, so the text is read from the shapes
    themselves -- this is also what proves the content is genuinely on the slide.
    """
    parts: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
            parts.append(shape.text_frame.text)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _deck_texts(deck) -> list[str]:
    return [_slide_text(slide) for slide in deck.slides]


# --------------------------------------------------------------------------- #
# The complete workflow
# --------------------------------------------------------------------------- #


def test_the_full_workflow_from_discussion_to_exported_deck(tmp_path: Path) -> None:
    """The whole scenario, in order, exactly as a user walks it."""
    assistant = ScriptedAssistant()
    service = _service(tmp_path, assistant=assistant)

    # 1. Create a new presentation project.
    project = service.new_project(
        "Reservoir Characterisation", audience="petroleum engineering students"
    )
    assert project.state == "DISCUSSION"
    assert service.list_ids() == [project.project_id]

    # 2. Discuss a technical subject with the assistant.
    discussion = service.discuss(
        project.project_id, "What controls porosity?", conversation_id="conv-e2e-1"
    )
    assert discussion.answer  # the assistant answered
    assert assistant.questions  # it really was asked
    assert discussion.project_id == project.project_id
    # The conversation reference is recorded on the project (never a copy of it).
    assert service.get(project.project_id).conversation_ids == ["conv-e2e-1"]

    # 3. Generate an outline with multiple sections and subsections.
    revision = service.draft_outline(project.project_id, text=OUTLINE)
    assert len(revision.sections) == 2
    assert revision.sections[0].heading == "Reservoir Characterisation"
    assert revision.sections[0].subsections == ("Porosity and storage", "Permeability and flow")
    assert revision.sections[1].heading == "Well Logs"
    assert service.get(project.project_id).state == "OUTLINE_DRAFT"

    # 4. Edit and approve the outline.
    edited = service.edit_outline(
        project.project_id,
        [
            OutlineSection("Reservoir Characterisation", ("Porosity", "Permeability"), 3),
            OutlineSection("Well Logs", ("Gamma ray", "Resistivity"), 2),
        ],
        note="tightened the headings",
    )
    assert edited.revision == revision.revision + 1
    assert edited.sections[0].estimated_slides == 3
    # Editing the outline moves it into review (the edit itself opens the review).
    assert service.get(project.project_id).state == "OUTLINE_REVIEW"
    service.approve_outline(project.project_id)
    assert service.get(project.project_id).state == "SLIDE_DRAFTING"
    assert service.get(project.project_id).approved_outline_revision == edited.revision

    # 5. Generate the first 5 slides (one batch).
    service.set_batch_size(5)
    first_batch = service.draft_slides_from_outline(project.project_id)
    assert 1 <= len(first_batch) <= 5
    assert [slide.number for slide in first_batch] == list(range(1, len(first_batch) + 1))

    # 6. Edit one slide and request a revision to another.
    target = first_batch[0]
    service.edit_slide(
        project.project_id, target.slide_id, title="Porosity: the storage term"
    )
    assert service.get(project.project_id).slide(target.slide_id).title == (
        "Porosity: the storage term"
    )
    if len(first_batch) > 1:
        revised = service.revise_slides(
            project.project_id,
            [first_batch[1].number],
            instruction="Add the units.",
        )
        assert revised and revised[0].slide_id == first_batch[1].slide_id
        # A revision is recorded so the change is auditable.
        assert service.revisions(project.project_id, first_batch[1].slide_id)

    # 7. Approve the batch and continue generating subsequent batches.
    service.approve_slides_in_range(
        project.project_id, 1, len(first_batch)
    )
    approved = {slide.number for slide in service.get(project.project_id).approved_slides}
    assert approved == {slide.number for slide in first_batch}
    while True:
        batch = service.draft_slides_from_outline(project.project_id)
        if not batch:
            break
        service.approve_slides_in_range(
            project.project_id, batch[0].number, batch[-1].number
        )
    project_state = service.get(project.project_id)
    assert len(project_state.slides) >= 4  # every outline topic produced a slide
    assert project_state.all_slides_approved

    # 8. Reorder two slides and verify the displayed numbering updates.
    service.review_slides(project.project_id)
    ordered_before = service.get(project.project_id).ordered_slides()
    first_id, second_id = ordered_before[0].slide_id, ordered_before[1].slide_id
    first_title, second_title = ordered_before[0].title, ordered_before[1].title
    service.reorder_slide(project.project_id, second_id, 1)
    ordered_after = service.get(project.project_id).ordered_slides()
    assert [slide.slide_id for slide in ordered_after[:2]] == [second_id, first_id]
    # Displayed numbering is dense and matches the new order.
    assert [slide.number for slide in ordered_after] == list(
        range(1, len(ordered_after) + 1)
    )
    assert service.get(project.project_id).slide_by_number(1).title == second_title
    assert service.get(project.project_id).slide_by_number(2).title == first_title

    # 9. Close and reopen the project; all approved content is preserved.
    reopened = PresentationService(
        root=tmp_path / "store", output_dir=tmp_path / "out", ask=ScriptedAssistant()
    )
    preserved = reopened.get(project.project_id)
    assert [slide.slide_id for slide in preserved.ordered_slides()] == [
        slide.slide_id for slide in ordered_after
    ]
    assert [slide.title for slide in preserved.ordered_slides()] == [
        slide.title for slide in ordered_after
    ]
    assert [slide.content for slide in preserved.ordered_slides()] == [
        slide.content for slide in ordered_after
    ]

    # 10. Approve the complete presentation.
    reopened.approve_slides(project.project_id)
    reopened.approve_content(project.project_id)
    final = reopened.get(project.project_id)
    assert final.state == STATE_CONTENT_APPROVED
    assert final.all_slides_approved

    # 11. Export the project to PowerPoint.
    report = reopened.export(project.project_id)
    assert report.ok, report.message
    assert Path(report.path).exists()
    assert Path(report.path).suffix == ".pptx"
    assert report.slide_count == len(final.slides)
    assert reopened.get(project.project_id).state == STATE_EXPORTED

    # 12. Reopen the .pptx programmatically and validate its structure.
    if PPTX_AVAILABLE:
        import pptx

        deck = pptx.Presentation(report.path)
        assert len(deck.slides) == len(final.slides)
        texts = _deck_texts(deck)
        # Each approved slide's title is on the matching slide, in the order the
        # user approved, so the written deck preserves the approved sequence.
        for index, slide in enumerate(final.ordered_slides()):
            assert slide.title in texts[index], (index, slide.title, texts[index])


@requires_pptx
def test_the_exported_deck_has_real_titles_content_formatting_and_notes(
    tmp_path: Path,
) -> None:
    """The written file is a real deck: shapes, formatting, and speaker notes."""
    import pptx
    from pptx.util import Pt

    service = _service(tmp_path)
    project = service.new_project("Formatting check")
    service.draft_outline(project.project_id, text="One\n- A\n- B")
    service.approve_outline(project.project_id)
    service.add_slides(
        project.project_id,
        [
            SlideSpec(
                title="Porosity",
                content=["Void fraction", "Connected pores"],
                notes="Remember to define effective porosity here.",
            ),
            SlideSpec(
                title="Permeability",
                content=["Transmits fluid", "Measured in darcies"],
            ),
        ],
    )
    service.review_slides(project.project_id)
    service.approve_slides(project.project_id)
    service.approve_content(project.project_id)
    report = service.export(project.project_id)
    assert report.ok, report.message

    deck = pptx.Presentation(report.path)
    assert len(deck.slides) == 2

    # Each slide carries its own title as real text.
    first_text = _slide_text(deck.slides[0])
    second_text = _slide_text(deck.slides[1])
    assert "Porosity" in first_text
    assert "Permeability" in second_text

    # Body content is present as real text, not an image.
    assert "Void fraction" in first_text
    assert "Connected pores" in first_text

    # Formatting: the heading run has a real point size and a colour.
    title_frame = next(
        shape.text_frame
        for shape in deck.slides[0].shapes
        if shape.has_text_frame and "Porosity" in shape.text_frame.text
    )
    title_run = title_frame.paragraphs[0].runs[0]
    assert title_run.font.size is not None
    assert title_run.font.size > Pt(20)
    assert title_run.font.color and title_run.font.color.rgb is not None

    # Speaker notes survive saving and reopening.
    assert deck.slides[0].has_notes_slide
    assert "effective porosity" in deck.slides[0].notes_slide.notes_text_frame.text
    assert not deck.slides[1].has_notes_slide or not (
        deck.slides[1].notes_slide.notes_text_frame.text.strip()
    )

    # The deck is a genuine OOXML package with the expected parts.
    with zipfile.ZipFile(report.path) as package:
        names = set(package.namelist())
    assert "ppt/presentation.xml" in names
    assert "[Content_Types].xml" in names
    assert sum(1 for name in names if name.startswith("ppt/slides/slide")) >= 2


# --------------------------------------------------------------------------- #
# The workflow is reachable through the CLI, too
# --------------------------------------------------------------------------- #


def test_the_cli_walks_the_workflow_and_exports(tmp_path: Path, capsys) -> None:
    """The headless twin drives the same service to a written deck."""
    from cli import presentation as cli_presentation

    store = tmp_path / "store"
    out = tmp_path / "out"

    def run(*argv: str) -> int:
        return cli_presentation.main(
            ["--root", str(store), "--output-dir", str(out), *argv]
        )

    assert run("new", "CLI Deck") == 0
    capsys.readouterr()
    assert run("--json", "list") == 0
    project_id = json.loads(capsys.readouterr().out)[0]["project_id"]

    assert run("outline", project_id, "--text", OUTLINE) == 0
    assert run("review-outline", project_id) == 0
    assert run("approve-outline", project_id) == 0
    assert run("add", project_id, "--title", "Porosity", "--bullet", "Void fraction") == 0
    assert run("add", project_id, "--title", "Permeability", "--bullet", "Flow") == 0
    assert run("review-slides", project_id) == 0
    assert run("approve-slides", project_id) == 0
    assert run("approve-content", project_id) == 0
    assert run("validate", project_id) == 0

    export_out = tmp_path / "deck.pptx"
    assert run("export", project_id, "--output", str(export_out)) == 0
    assert export_out.exists()
    assert export_out.suffix == ".pptx"

    # The history records the successful export.
    assert run("history", project_id) == 0
    history = capsys.readouterr().out
    assert "EXP-001" in history or "Exported" in history


def test_the_cli_demo_builds_and_exports_a_sample_deck(tmp_path: Path, capsys) -> None:
    """`python -m cli presentation demo` is the one-command sample."""
    from cli import presentation as cli_presentation

    out = tmp_path / "demo.pptx"
    code = cli_presentation.main(
        [
            "--root",
            str(tmp_path / "store"),
            "--output-dir",
            str(tmp_path / "out"),
            "demo",
            "--title",
            "Sample",
            "--output",
            str(out),
        ]
    )
    captured = capsys.readouterr().out
    assert code == 0, captured
    assert out.exists()
    assert "Built project" in captured


# --------------------------------------------------------------------------- #
# Internationalisation, equations, tables, and long titles
# --------------------------------------------------------------------------- #


@requires_pptx
def test_arabic_english_mixed_equations_tables_and_long_titles_survive(
    tmp_path: Path,
) -> None:
    """Unicode, equations, tables, and a long title all round-trip through export."""
    import pptx

    service = _service(tmp_path)
    project = service.new_project("ثنائي اللغة / Bilingual")
    service.draft_outline(project.project_id, text="Bilingual\n- Arabic\n- English")
    service.approve_outline(project.project_id)

    long_title = (
        "A Comprehensive Review of Porosity, Permeability, and Fluid Saturation "
        "Measurements in Heterogeneous Carbonate Reservoirs"
    )
    service.add_slides(
        project.project_id,
        [
            SlideSpec(title=long_title, content=["Long heading fits the slide"]),
            SlideSpec(title="الخصائص البتروفيزيائية", content=["المسامية", "النفاذية"]),
            SlideSpec(
                title="Mixed direction",
                content=["Porosity نسبة المسامية", "Permeability النفاذية"],
            ),
            SlideSpec(
                title="Darcy's law",
                content=[],
                content_type=CONTENT_EQUATION,
                equations=["Q = k A ΔP / (μ L)"],
            ),
            SlideSpec(
                title="Core measurements",
                content=[],
                content_type=CONTENT_TABLE,
                table=[
                    ["Sample", "Porosity %", "Permeability mD"],
                    ["A-1", "18.2", "145"],
                    ["A-2", "21.0", "230"],
                ],
            ),
        ],
    )
    service.review_slides(project.project_id)
    service.approve_slides(project.project_id)
    service.approve_content(project.project_id)
    report = service.export(project.project_id, settings=service.get(project.project_id).settings.replace(rtl=True))
    assert report.ok, report.message

    deck = pptx.Presentation(report.path)
    assert len(deck.slides) == 5
    texts = _deck_texts(deck)
    assert long_title.split(",")[0] in texts[0]
    assert "المسامية" in texts[1]
    assert "نسبة المسامية" in texts[2]
    assert "Q = k A" in texts[3]
    assert "18.2" in texts[4]

    # The table is a real table shape, not a text block.
    tables = [shape for shape in deck.slides[4].shapes if shape.has_table]
    assert tables, "the table slide carries no table shape"
    table = tables[0].table
    assert len(table.rows) == 3
    assert table.cell(0, 1).text == "Porosity %"
    assert table.cell(2, 2).text == "230"

    # The long title did not overflow the slide width (no shape is wider than it).
    for slide in deck.slides:
        for shape in slide.shapes:
            assert shape.left >= 0
            assert shape.left + shape.width <= deck.slide_width + 1


# --------------------------------------------------------------------------- #
# Unapproved or outdated content cannot be exported
# --------------------------------------------------------------------------- #


def test_unapproved_content_cannot_be_exported(tmp_path: Path) -> None:
    """The export gate refuses a deck that is not fully approved."""
    from core.presentation import PresentationError

    service = _service(tmp_path)
    project = service.new_project("Gate")
    service.draft_outline(project.project_id, text="One\n- A\n- B")
    service.approve_outline(project.project_id)
    service.add_slide(project.project_id, SlideSpec(title="Only", content=["x"]))
    service.review_slides(project.project_id)

    allowed, reason = service.can_export(project.project_id)
    assert not allowed and reason
    with pytest.raises(PresentationError):
        service.export(project.project_id)

    # Approve the slide but not the content: still refused.
    slide = service.get(project.project_id).ordered_slides()[0]
    service.approve_slide(project.project_id, slide.slide_id)
    assert not service.can_export(project.project_id)[0]

    # Approving the content opens the gate.
    service.approve_content(project.project_id)
    assert service.can_export(project.project_id)[0]


def test_content_changed_after_approval_is_outdated_and_refused(tmp_path: Path) -> None:
    """Editing an approved slide re-arms the gate until it is approved again."""
    from core.presentation import PresentationError

    service = _service(tmp_path)
    project = service.new_project("Outdated")
    service.draft_outline(project.project_id, text="One\n- A\n- B")
    service.approve_outline(project.project_id)
    service.add_slides(
        project.project_id,
        [
            SlideSpec(title="A", content=["a"]),
            SlideSpec(title="B", content=["b"]),
        ],
    )
    service.review_slides(project.project_id)
    service.approve_slides(project.project_id)
    service.approve_content(project.project_id)
    assert service.can_export(project.project_id)[0]

    # A change after approval marks the project's approval outdated.
    slide = service.get(project.project_id).slide_by_number(2)
    service.edit_slide(project.project_id, slide.slide_id, content=["b", "b2"])
    state = service.get(project.project_id)
    assert state.state == STATE_SLIDE_REVIEW
    assert state.content_approved is False
    allowed, reason = service.can_export(project.project_id)
    assert not allowed
    with pytest.raises(PresentationError):
        service.export(project.project_id)

    # Re-approving the changed slide and the content restores the gate.
    service.approve_slide(project.project_id, slide.slide_id)
    service.approve_content(project.project_id)
    assert service.can_export(project.project_id)[0]


# --------------------------------------------------------------------------- #
# A crash during the write must never strand the project
# --------------------------------------------------------------------------- #


def test_an_interrupted_export_is_recovered_not_stranded(tmp_path: Path) -> None:
    """A project left mid-export (a crash) is recoverable, never a dead end.

    ``EXPORTING`` has no outgoing edge except to ``EXPORTED``/``EXPORT_FAILED``,
    so a project interrupted *between* the state write and the finished file
    would otherwise refuse every export, edit, and reopen forever. Loading it
    returns it to ``CONTENT_APPROVED`` so the export can simply be run again.
    """
    from core.presentation import STATE_EXPORTING

    service = _service(tmp_path)
    project = service.new_project("Interrupted")
    service.draft_outline(project.project_id, text="One\n- A\n- B")
    service.approve_outline(project.project_id)
    service.add_slide(project.project_id, SlideSpec(title="Only", content=["x"]))
    service.approve_slides(project.project_id)
    service.approve_content(project.project_id)

    # Simulate a hard interruption: the state is persisted as EXPORTING.
    stranded = service.get(project.project_id)
    stranded.state = STATE_EXPORTING
    service._save(stranded)
    on_disk = service._store.load(project.project_id)
    assert on_disk is not None and on_disk.state == STATE_EXPORTING

    # A fresh service (as after a restart) recovers it on load.
    reopened = PresentationService(
        root=tmp_path / "store", output_dir=tmp_path / "out", ask=ScriptedAssistant()
    )
    recovered = reopened.get(project.project_id)
    assert recovered.state == STATE_CONTENT_APPROVED
    assert reopened.can_export(project.project_id)[0]

    # The export runs, and the project ends in EXPORTED as normal.
    report = reopened.export(project.project_id)
    assert report.ok, report.message
    assert Path(report.path).exists()
    assert reopened.get(project.project_id).state == STATE_EXPORTED

    # Editing after an interruption also works (the second recovery path).
    again = reopened.get(project.project_id)
    again.state = STATE_EXPORTING
    reopened._save(again)
    slide = reopened.get(project.project_id).slide_by_number(1)
    reopened.edit_slide(project.project_id, slide.slide_id, title="Edited after crash")
    assert reopened.get(project.project_id).slide_by_number(1).title == (
        "Edited after crash"
    )


def test_the_recovery_leaves_an_exported_project_alone(tmp_path: Path) -> None:
    """The recovery only touches EXPORTING; an exported deck is untouched."""
    service = _service(tmp_path)
    project = service.new_project("Exported")
    service.draft_outline(project.project_id, text="One\n- A\n- B")
    service.approve_outline(project.project_id)
    service.add_slide(project.project_id, SlideSpec(title="Only", content=["x"]))
    service.approve_slides(project.project_id)
    service.approve_content(project.project_id)
    first = service.export(project.project_id)
    assert first.ok

    # Re-loading an EXPORTED project keeps it EXPORTED and still exportable.
    reloaded = service.get(project.project_id)
    assert reloaded.state == STATE_EXPORTED
    assert service.can_export(project.project_id)[0]
    second = service.export(project.project_id)
    assert second.ok
    assert Path(first.path).exists() and Path(second.path).exists()


# --------------------------------------------------------------------------- #
# The real assistant (grounded, offline) drives the discussion
# --------------------------------------------------------------------------- #


def test_the_real_grounded_assistant_answers_a_technical_question(
    tmp_path: Path,
) -> None:
    """The project's own assistant answers, and its answer becomes slide content.

    This wires the *real* ``ai.jarvis.build_session(text_only=True)`` session
    (the grounded assistant + the in-project retriever) rather than a stub, so
    the discussion path is proven against the assistant the product ships. It is
    offline and deterministic, so it needs no model and no network.
    """
    from ai.jarvis import build_session

    session = build_session(text_only=True)

    def ask(question: str) -> str:
        turn = session.handle(question)
        return turn.answer_text or turn.response

    service = _service(tmp_path, assistant=ask)
    project = service.new_project("Real assistant")

    # 1. Discuss a real petroleum-engineering question.
    result = service.discuss(
        project.project_id, "What controls porosity and permeability in a reservoir?"
    )
    assert result.answer.strip(), "the grounded assistant returned no answer"

    # 2. The answer is real prose that parses into an outline.
    revision = service.draft_outline(project.project_id, text=result.answer)
    assert revision.sections, "the assistant's answer produced no outline sections"

    # 3. The whole pipeline still runs, and the deck is written.
    service.approve_outline(project.project_id)
    slides = service.draft_slides_from_outline(project.project_id)
    assert slides, "drafting produced no slides from the assistant's answer"
    service.review_slides(project.project_id)
    service.approve_slides(project.project_id)
    service.approve_content(project.project_id)
    report = service.export(project.project_id)
    assert report.ok, report.message
    assert Path(report.path).exists()
    assert report.slide_count == len(service.get(project.project_id).slides)


@requires_pptx
def test_an_image_slide_embeds_a_real_picture(tmp_path: Path) -> None:
    """An image slide carries a real picture shape, and a bad path is flagged."""
    import pptx

    from core.imagegen import encode_png

    png = tmp_path / "figure.png"
    png.write_bytes(
        encode_png(
            20,
            20,
            [bytes([(row * 10) % 256, 100, 200] * 20) for row in range(20)],
        )
    )

    service = _service(tmp_path)
    project = service.new_project("With a figure")
    service.draft_outline(project.project_id, text="One\n    - A\n    - B")
    service.approve_outline(project.project_id)
    service.add_slide(
        project.project_id,
        SlideSpec(
            title="A figure",
            content=["See the plot"],
            content_type=CONTENT_IMAGE,
            images=[str(png)],
        ),
    )
    service.review_slides(project.project_id)
    service.approve_slides(project.project_id)
    service.approve_content(project.project_id)
    report = service.export(project.project_id)
    assert report.ok, report.message

    deck = pptx.Presentation(report.path)
    pictures = [shape for shape in deck.slides[0].shapes if shape.shape_type == 13]
    assert pictures, "the image slide carries no picture shape"
    assert pictures[0].width > 0 and pictures[0].height > 0


def test_a_missing_image_is_reported_not_silently_dropped(tmp_path: Path) -> None:
    """A referenced image that does not exist is an advisory warning."""
    service = _service(tmp_path)
    project = service.new_project("Missing figure")
    service.draft_outline(project.project_id, text="One\n    - A\n    - B")
    service.approve_outline(project.project_id)
    slide = service.add_slide(
        project.project_id,
        SlideSpec(
            title="Broken figure",
            content=["text"],
            content_type=CONTENT_IMAGE,
            images=[str(tmp_path / "does-not-exist.png")],
        ),
    )
    # The slide itself is valid (it has text); the missing image is advisory.
    service.approve_slide(project.project_id, slide.slide_id)
    warnings = service.export_warnings(project.project_id)
    assert warnings, "a missing image produced no warning"
    assert any("image" in warning.lower() and "not found" in warning.lower() for warning in warnings)


def test_an_empty_slide_cannot_be_approved_or_exported(tmp_path: Path) -> None:
    """A slide with no content is refused, so the deck is never hollow."""
    from core.presentation import PresentationError

    service = _service(tmp_path)
    project = service.new_project("Empty")
    service.draft_outline(project.project_id, text="One\n- A\n- B")
    service.approve_outline(project.project_id)
    slide = service.add_slide(project.project_id, SlideSpec(title="", content=[]))
    with pytest.raises(PresentationError):
        service.approve_slide(project.project_id, slide.slide_id)
    with pytest.raises(PresentationError):
        service.approve_content(project.project_id)
